"""Gera o DOCX do Texto para Discussao a partir do Markdown revisado."""

from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "artigo" / "TD_clusterizacao_noticias_pf.md"
OUTPUT = ROOT / "artigo" / "TD_clusterizacao_noticias_pf.docx"
BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
MUTED = RGBColor(95, 102, 112)
TABLE_FILL = "F4F6F9"
TABLE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120


def font(run, name="Calibri", size=11, color=None, bold=None, italic=None):
    run.font.name = name
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_spacing(style, before, after, line_spacing):
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.line_spacing = line_spacing


def configure_document(doc):
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    set_spacing(normal, 0, 8, 1.333)

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 12, 6),
        ("Heading 3", 12, DARK_BLUE, 8, 4),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.font.bold = True
        set_spacing(style, before, after, 1.15)

    for name in ("List Number", "List Bullet"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.194)
        set_spacing(style, 0, 4, 1.208)

    code = doc.styles["Normal"].base_style
    if "Code Block" not in [s.name for s in doc.styles]:
        code = doc.styles.add_style("Code Block", 1)
    code.font.name = "Consolas"
    code._element.rPr.rFonts.set(qn("w:ascii"), "Consolas")
    code._element.rPr.rFonts.set(qn("w:hAnsi"), "Consolas")
    code.font.size = Pt(8.5)
    code.paragraph_format.left_indent = Inches(0.16)
    code.paragraph_format.right_indent = Inches(0.12)
    set_spacing(code, 3, 7, 1.05)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("Texto para Discussão | Metodologia incremental em bases textuais")
    font(run, size=8.5, color=MUTED)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_before = Pt(0)
    run = footer.add_run("Página ")
    font(run, size=9, color=MUTED)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)


def add_hyperlink(paragraph, label, target):
    relationship = paragraph.part.relate_to(
        target,
        "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
        is_external=True,
    )
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), relationship)
    run = OxmlElement("w:r")
    props = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), "0563C1")
    props.append(color)
    underline = OxmlElement("w:u")
    underline.set(qn("w:val"), "single")
    props.append(underline)
    run.append(props)
    text = OxmlElement("w:t")
    text.text = label
    run.append(text)
    link.append(run)
    paragraph._p.append(link)


def add_inline(paragraph, text):
    token_pattern = re.compile(r"(\[[^\]]+\]\(https?://[^)]+\)|`[^`]+`|\*[^*]+\*)")
    for piece in token_pattern.split(text):
        if not piece:
            continue
        link = re.fullmatch(r"\[([^\]]+)\]\((https?://[^)]+)\)", piece)
        if link:
            add_hyperlink(paragraph, link.group(1), link.group(2))
        elif piece.startswith("`") and piece.endswith("`"):
            run = paragraph.add_run(piece[1:-1])
            font(run, name="Consolas", size=9)
        elif piece.startswith("*") and piece.endswith("*"):
            run = paragraph.add_run(piece[1:-1])
            font(run, italic=True)
        else:
            paragraph.add_run(piece)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{edge}"))
        if node is None:
            node = OxmlElement(f"w:{edge}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def table_widths(column_count):
    choices = {
        2: [2700, 6660],
        3: [2000, 3500, 3860],
        4: [1800, 2500, 2500, 2560],
        5: [1600, 1940, 1940, 1940, 1940],
        6: [1550, 1550, 1550, 1550, 1550, 1610],
    }
    return choices.get(column_count, [TABLE_WIDTH_DXA // column_count] * column_count)


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    width = tbl_pr.first_child_found_in("w:tblW")
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(TABLE_WIDTH_DXA))
    indent = tbl_pr.first_child_found_in("w:tblInd")
    if indent is None:
        indent = OxmlElement("w:tblInd")
        tbl_pr.append(indent)
    indent.set(qn("w:type"), "dxa")
    indent.set(qn("w:w"), str(TABLE_INDENT_DXA))
    grid = table._tbl.tblGrid
    while len(grid):
        grid.remove(grid[0])
    for width_value in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width_value))
        grid.append(col)
    for row in table.rows:
        for cell, width_value in zip(row.cells, widths):
            cell.width = Inches(width_value / 1440)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tc_width = cell._tc.get_or_add_tcPr().tcW
            tc_width.set(qn("w:type"), "dxa")
            tc_width.set(qn("w:w"), str(width_value))
            set_cell_margins(cell)


def add_table(doc, rows):
    if not rows:
        return
    column_count = len(rows[0])
    table = doc.add_table(rows=0, cols=column_count)
    table.style = "Table Grid"
    widths = table_widths(column_count)
    for row_index, values in enumerate(rows):
        cells = table.add_row().cells
        if row_index == 0:
            tr_pr = table.rows[0]._tr.get_or_add_trPr()
            header = OxmlElement("w:tblHeader")
            header.set(qn("w:val"), "true")
            tr_pr.append(header)
        for col_index, value in enumerate(values):
            paragraph = cells[col_index].paragraphs[0]
            paragraph.paragraph_format.space_before = Pt(0)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.line_spacing = 1.1
            add_inline(paragraph, value.strip())
            for run in paragraph.runs:
                font(run, size=9.2, bold=row_index == 0)
            if row_index == 0:
                shading = OxmlElement("w:shd")
                shading.set(qn("w:fill"), TABLE_FILL)
                cells[col_index]._tc.get_or_add_tcPr().append(shading)
    set_table_geometry(table, widths)
    after = doc.add_paragraph()
    after.paragraph_format.space_after = Pt(2)


def add_cover(doc, title):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(42)
    p.paragraph_format.space_after = Pt(12)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("TEXTO PARA DISCUSSÃO")
    font(run, size=10, color=BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(14)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run(title)
    font(run, size=25, color=DARK_BLUE, bold=True)

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(28)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run("Clusterização, classificação auditável e aprendizado contínuo")
    font(run, size=12.5, color=MUTED, italic=True)


def add_image(doc, alt, relative_path):
    path = SOURCE.parent / relative_path
    if not path.exists():
        paragraph = doc.add_paragraph(f"[Figura ausente: {alt}]")
        paragraph.paragraph_format.space_after = Pt(6)
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(8)
    paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run()
    shape = run.add_picture(str(path), width=Inches(6.1))
    shape._inline.docPr.set("descr", alt)
    caption = doc.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(9)
    run = caption.add_run(f"Figura. {alt}.")
    font(run, size=9, color=MUTED, italic=True)


def render_markdown(doc, lines):
    index = 0
    first_heading = True
    paragraph_buffer = []

    def flush_paragraph():
        if paragraph_buffer:
            paragraph = doc.add_paragraph()
            add_inline(paragraph, " ".join(part.strip() for part in paragraph_buffer))
            paragraph_buffer.clear()

    while index < len(lines):
        line = lines[index].rstrip()
        if not line:
            flush_paragraph()
            index += 1
            continue
        if line.startswith("```"):
            flush_paragraph()
            code_lines = []
            index += 1
            while index < len(lines) and not lines[index].startswith("```"):
                code_lines.append(lines[index].rstrip())
                index += 1
            paragraph = doc.add_paragraph(style="Code Block")
            run = paragraph.add_run("\n".join(code_lines))
            font(run, name="Consolas", size=8.5)
            index += 1
            continue
        heading = re.match(r"^(#{1,3})\s+(.*)$", line)
        if heading:
            flush_paragraph()
            level, title = len(heading.group(1)), heading.group(2)
            if level == 1 and first_heading:
                add_cover(doc, title)
                first_heading = False
            else:
                paragraph = doc.add_paragraph(style=f"Heading {max(1, level - 1)}")
                add_inline(paragraph, title)
            index += 1
            continue
        image = re.match(r"^!\[([^\]]+)\]\(([^)]+)\)$", line)
        if image:
            flush_paragraph()
            add_image(doc, image.group(1), image.group(2))
            index += 1
            continue
        if line.startswith("|") and line.endswith("|"):
            flush_paragraph()
            rows = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                values = [value.strip() for value in lines[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", value) for value in values):
                    rows.append(values)
                index += 1
            add_table(doc, rows)
            continue
        numbered = re.match(r"^\d+\.\s+(.*)$", line)
        bullet = re.match(r"^-\s+(.*)$", line)
        if numbered or bullet:
            flush_paragraph()
            paragraph = doc.add_paragraph(style="List Number" if numbered else "List Bullet")
            add_inline(paragraph, (numbered or bullet).group(1))
            index += 1
            continue
        paragraph_buffer.append(line)
        index += 1
    flush_paragraph()


def main():
    markdown = SOURCE.read_text(encoding="utf-8").splitlines()
    doc = Document()
    configure_document(doc)
    render_markdown(doc, markdown)
    core = doc.core_properties
    core.title = "Uma Metodologia Incremental para Clusterização, Classificação e Aprendizado Contínuo em Grandes Bases Textuais"
    core.subject = "Texto para Discussão"
    core.author = "Projeto NT PF"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
