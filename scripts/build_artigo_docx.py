from __future__ import annotations

import re
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
MD_PATH = ROOT / "artigo" / "TD_clusterizacao_noticias_pf.md"
OUT_PATH = ROOT / "artigo" / "TD_clusterizacao_noticias_pf_revisado.docx"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BODY_FONT = "Calibri"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_width(cell, width_dxa: int) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.find(qn("w:tcW"))
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths: list[int]) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.find(qn("w:tblInd"))
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    tbl_layout = tbl_pr.find(qn("w:tblLayout"))
    if tbl_layout is None:
        tbl_layout = OxmlElement("w:tblLayout")
        tbl_pr.append(tbl_layout)
    tbl_layout.set(qn("w:type"), "fixed")

    tbl_grid = table._tbl.tblGrid
    if tbl_grid is None:
        tbl_grid = OxmlElement("w:tblGrid")
        table._tbl.insert(1, tbl_grid)
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            set_cell_width(cell, widths[min(idx, len(widths) - 1)])
            cell.vertical_alignment = WD_ALIGN_VERTICAL.TOP


def set_cell_margins(table, top=80, start=120, bottom=80, end=120) -> None:
    tbl_pr = table._tbl.tblPr
    tbl_cell_mar = tbl_pr.find(qn("w:tblCellMar"))
    if tbl_cell_mar is None:
        tbl_cell_mar = OxmlElement("w:tblCellMar")
        tbl_pr.append(tbl_cell_mar)
    for m, v in {"top": top, "start": start, "bottom": bottom, "end": end}.items():
        node = tbl_cell_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tbl_cell_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_run_font(run, size: float | None = None, bold: bool | None = None, italic: bool | None = None, color: str | None = None, font: str = BODY_FONT) -> None:
    run.font.name = font
    run._element.rPr.rFonts.set(qn("w:eastAsia"), font)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


INLINE_RE = re.compile(r"(`[^`]+`|\*\*[^*]+\*\*|\*[^*]+\*|\[[^\]]+\]\([^)]+\))")


def add_inline_text(paragraph, text: str, *, base_size: float = 11, italic_default: bool = False) -> None:
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("`") and part.endswith("`"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, base_size, font="Consolas")
        elif part.startswith("**") and part.endswith("**"):
            run = paragraph.add_run(part[2:-2])
            set_run_font(run, base_size, bold=True, italic=italic_default)
        elif part.startswith("*") and part.endswith("*"):
            run = paragraph.add_run(part[1:-1])
            set_run_font(run, base_size, italic=True)
        elif part.startswith("[") and "](" in part and part.endswith(")"):
            label, url = re.match(r"\[([^\]]+)\]\(([^)]+)\)", part).groups()
            run = paragraph.add_run(f"{label} ({url})")
            set_run_font(run, base_size, color="0563C1")
        else:
            run = paragraph.add_run(part)
            set_run_font(run, base_size, italic=italic_default)


def add_paragraph(doc: Document, text: str, style: str = "Normal", *, italic: bool = False) -> None:
    p = doc.add_paragraph(style=style)
    add_inline_text(p, text, italic_default=italic)


def split_table_row(line: str) -> list[str]:
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|"):
        line = line[:-1]
    return [cell.strip() for cell in line.split("|")]


def is_separator_row(cells: Iterable[str]) -> bool:
    return all(re.fullmatch(r":?-{3,}:?", cell.strip()) for cell in cells)


def table_widths(headers: list[str], rows: list[list[str]]) -> list[int]:
    cols = len(headers)
    samples = headers + [cell for row in rows[:8] for cell in row]
    if cols == 2:
        if any(len(x) > 80 for x in samples):
            return [2600, PAGE_WIDTH_DXA - 2600]
        return [3100, PAGE_WIDTH_DXA - 3100]
    if cols == 3:
        return [2100, 3600, PAGE_WIDTH_DXA - 2100 - 3600]
    if cols == 4:
        return [2200, 1800, 2600, PAGE_WIDTH_DXA - 6600]
    if cols == 5:
        return [1700, 1650, 2300, 1900, PAGE_WIDTH_DXA - 7550]
    if cols == 7:
        return [2800, 900, 1000, 1200, 1400, 1200, PAGE_WIDTH_DXA - 8500]
    if cols == 8:
        return [1500, 1000, 1000, 1300, 1250, 900, 1300, PAGE_WIDTH_DXA - 8250]
    base = PAGE_WIDTH_DXA // cols
    widths = [base] * cols
    widths[-1] += PAGE_WIDTH_DXA - sum(widths)
    return widths


def add_table(doc: Document, rows: list[list[str]]) -> None:
    headers = rows[0]
    body = rows[1:]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    widths = table_widths(headers, body)
    set_table_geometry(table, widths)
    set_cell_margins(table)

    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]
        cell.text = ""
        p = cell.paragraphs[0]
        add_inline_text(p, header, base_size=8.2)
        for run in p.runs:
            run.bold = True
        set_cell_shading(cell, "F2F4F7")

    for row_values in body:
        row = table.add_row()
        for idx, value in enumerate(row_values):
            cell = row.cells[idx]
            cell.text = ""
            p = cell.paragraphs[0]
            add_inline_text(p, value, base_size=7.8 if len(headers) >= 5 else 8.5)

    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.line_spacing = 1.0


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    title = doc.styles["Title"]
    title.font.name = BODY_FONT
    title._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
    title.font.size = Pt(18)
    title.font.bold = True
    title.font.color.rgb = RGBColor.from_string("0B2545")
    title.paragraph_format.space_after = Pt(12)

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = doc.styles[name]
        style.font.name = BODY_FONT
        style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for style_name in ("Caption", "Intense Quote"):
        if style_name in doc.styles:
            style = doc.styles[style_name]
            style.font.name = BODY_FONT
            style._element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)
            style.font.size = Pt(9)
            style.font.italic = True
            style.font.color.rgb = RGBColor.from_string("555555")


def add_code_block(doc: Document, code: str) -> None:
    for line in code.rstrip("\n").splitlines():
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.space_after = Pt(0)
        run = p.add_run(line if line else " ")
        set_run_font(run, 8.5, font="Consolas")


def add_image(doc: Document, md_dir: Path, line: str) -> None:
    match = re.match(r"!\[([^\]]*)\]\(([^)]+)\)", line.strip())
    if not match:
        return
    alt, src = match.groups()
    image_path = (md_dir / src).resolve()
    if not image_path.exists():
        add_paragraph(doc, f"[Imagem não encontrada: {src}]", italic=True)
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = p.add_run()
    run.add_picture(str(image_path), width=Inches(6.35))


def should_landscape(rows: list[list[str]]) -> bool:
    return len(rows[0]) >= 7


def build_docx() -> None:
    text = MD_PATH.read_text(encoding="utf-8")
    lines = text.splitlines()
    md_dir = MD_PATH.parent
    doc = Document()
    style_document(doc)

    first_title = True
    i = 0
    in_code = False
    code_lines: list[str] = []

    while i < len(lines):
        line = lines[i]

        if line.strip().startswith("```"):
            if in_code:
                add_code_block(doc, "\n".join(code_lines))
                code_lines = []
                in_code = False
            else:
                in_code = True
            i += 1
            continue

        if in_code:
            code_lines.append(line)
            i += 1
            continue

        stripped = line.strip()
        if not stripped:
            i += 1
            continue

        if stripped.startswith("|"):
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = split_table_row(lines[i])
                if not is_separator_row(cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue

        if stripped.startswith("!["):
            add_image(doc, md_dir, stripped)
            i += 1
            continue

        if stripped.startswith("# "):
            p = doc.add_paragraph(style="Title")
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            add_inline_text(p, stripped[2:].strip(), base_size=18)
            first_title = False
        elif stripped.startswith("## "):
            if not first_title and stripped != "## Resumo":
                pass
            add_paragraph(doc, stripped[3:].strip(), style="Heading 1")
        elif stripped.startswith("### "):
            add_paragraph(doc, stripped[4:].strip(), style="Heading 2")
        elif stripped.startswith("**") and stripped.endswith("**"):
            add_paragraph(doc, stripped[2:-2], style="Caption")
        elif stripped.startswith("*") and stripped.endswith("*"):
            add_paragraph(doc, stripped[1:-1], style="Caption", italic=True)
        elif re.match(r"^\d+\.\s+", stripped):
            p = doc.add_paragraph(style="List Number")
            add_inline_text(p, re.sub(r"^\d+\.\s+", "", stripped))
        elif stripped.startswith("- "):
            p = doc.add_paragraph(style="List Bullet")
            add_inline_text(p, stripped[2:])
        else:
            add_paragraph(doc, stripped)
        i += 1

    OUT_PATH.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUT_PATH)
    print(OUT_PATH)


if __name__ == "__main__":
    build_docx()
